# -*- coding: utf-8 -*-
"""40_verify_pptx.py —— 校验 PPTX 与 PDF/浏览器是否一致（成品验收）。

    python 40_verify_pptx.py --pptx poster.pptx --shapes _work/shapes.json \
        --html poster.html --pdf-size 900x1200 \
        --pdf-preview _work/preview.png --pptx-pdf _work/pptx.pdf \
        [--tol-center 1.0] [--tol-top 10.0]

五项：
  1) 编号居中：PPT 渲染图 vs HTML 预览图，徽标内数字的墨迹中心（≤tol-center）
  2) 文本落位：每个文本项在两图中的墨迹顶边差（中位/最大，≤tol-top）
  3) 图片：普通图不许变形；object-fit:cover 必须是**裁切**、contain 必须等比（不许拉伸）
  4) 包结构：内建 ZIP/XML/关系完整性检查
  5) PPTX→PDF 单页且尺寸符合目标
"""
from __future__ import annotations

import argparse
import json
import os
import sys
import xml.etree.ElementTree as ET
import zipfile

import numpy as np
from PIL import Image

REQUIRED_PARTS = ("[Content_Types].xml", "_rels/.rels", "ppt/presentation.xml",
                  "ppt/_rels/presentation.xml.rels", "ppt/slides/slide1.xml",
                  "ppt/slides/_rels/slide1.xml.rels")
for _s in (sys.stdout, sys.stderr):        # Windows 中文控制台默认 GBK：打印 "−/≈/✓" 会崩
    try:
        _s.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

PICTURE = 13


def ink_centroid(arr, s, box, diff_thr=40):
    """墨迹像素重心（canvas px）——比包围盒中心更抗边缘/抗锯齿干扰。"""
    x0, y0 = int(box["x"] * s), int(box["y"] * s)
    x1, y1 = int((box["x"] + box["w"]) * s), int((box["y"] + box["h"]) * s)
    sub = arr[max(0, y0):y1, max(0, x0):x1]
    if sub.size == 0 or sub.shape[0] < 3 or sub.shape[1] < 3:
        return None
    bg = np.median(sub.reshape(-1, sub.shape[2]), axis=0)
    m = np.abs(sub - bg).max(axis=2) > diff_thr
    if not m.any():
        return None
    ys, xs = np.where(m)
    return (x0 + xs.mean()) / s, (y0 + ys.mean()) / s


def ink_box(arr, s, box, pad_y=2, pad_x=2, diff_thr=40):
    """窗口内墨迹包围盒：与**窗口内背景中位数**差异 > diff_thr 才算墨迹。

    不要用绝对亮度阈值：面板标题条(#E8F1FD)、note 底色(#F2F7FE)、指标卡底色都接近白，
    绝对阈值会把底色算成墨迹，结论完全反（这套算法也能正确测深色底上的白字）。
    """
    x0, y0 = int((box["x"] + pad_x) * s), int((box["y"] - pad_y) * s)
    x1, y1 = int((box["x"] + box["w"] - pad_x) * s), int((box["y"] + box["h"] + pad_y) * s)
    sub = arr[max(0, y0):y1, max(0, x0):x1].astype(int)
    if sub.size == 0 or sub.shape[0] < 3 or sub.shape[1] < 3:
        return None
    bg = np.median(sub.reshape(-1, sub.shape[2]), axis=0)
    m = np.abs(sub - bg).max(axis=2) > diff_thr
    if not m.any():
        return None
    ys, xs = np.where(m)
    return (x0 + xs.min(), y0 + ys.min(), x0 + xs.max(), y0 + ys.max())


def check_package(path: str):
    """ZIP 完整性 + 必备部件 + 所有 XML/rels 可解析 + 内部关系目标存在。"""
    problems = []
    try:
        z = zipfile.ZipFile(path)
    except Exception as exc:
        return ["不是合法 zip：%s" % exc]
    with z:
        bad = z.testzip()
        if bad:
            problems.append("zip 内容损坏：%s" % bad)
        names = set(z.namelist())
        for req in REQUIRED_PARTS:
            if req not in names:
                problems.append("缺部件 %s" % req)
        for n in sorted(names):
            if not (n.endswith(".xml") or n.endswith(".rels")):
                continue
            try:
                root = ET.fromstring(z.read(n))
            except Exception as exc:
                problems.append("XML 解析失败 %s：%s" % (n, exc))
                continue
            if not n.endswith(".rels"):
                continue
            base = os.path.dirname(os.path.dirname(n))         # ppt/slides/_rels/x.rels → ppt/slides
            for rel in root:
                if rel.get("TargetMode") == "External":
                    continue
                full = os.path.normpath(os.path.join(base, (rel.get("Target") or "").lstrip("/")))
                if full.replace("\\", "/") not in names:
                    problems.append("关系目标缺失：%s -> %s" % (n, full))
    return problems


def parse_size(txt, flag="--pdf-size"):
    s = str(txt).lower().replace("×", "x").replace(",", "x").replace(" ", "")
    parts = s.split("x")
    if len(parts) != 2:
        raise ValueError("%s 要写成 WxH（例如 900x1200）；收到：%r" % (flag, txt))
    return float(parts[0]), float(parts[1])


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--pptx", required=True)
    ap.add_argument("--shapes", required=True)
    ap.add_argument("--pdf-preview", required=True, help="HTML 渲染的预览 PNG")
    ap.add_argument("--pptx-pdf", required=True, help="PPTX 转出的 PDF（用于栅格化）")
    ap.add_argument("--html", default=None, help="墙报 HTML —— 用于解析图片相对路径")
    ap.add_argument("--pdf-size", default=None, help="目标物理尺寸 mm，如 900x1200")
    ap.add_argument("--aspect-tol", type=float, default=0.5, help="图片纵横比容差 %%")
    ap.add_argument("--tol-center", type=float, default=1.0)
    ap.add_argument("--tol-top", type=float, default=10.0)
    ap.add_argument("--no-badge-check", action="store_true",
                    help="这份墙报确实没有编号徽标（hv 文本项）时显式跳过；默认「没有可检查的编号」= FAIL")
    ap.add_argument("--allow-missing-pictures", action="store_true",
                    help="PPTX 里明显少于实测图片数时也放行（默认 FAIL：缺图不算验收通过）")
    ap.add_argument("--canvas-w", type=float, default=None,
                    help="预览 PNG 对应的画布宽 px；缺省取 shapes.json 的 canvas.w")
    ap.add_argument("--dpi", type=float, default=150.0)
    ap.add_argument("--work", default="_work")
    args = ap.parse_args()

    for p in (args.pptx, args.shapes, args.pdf_preview, args.pptx_pdf):
        if not os.path.isfile(p):
            print("缺输入文件：%s" % p); return 2
    os.makedirs(args.work, exist_ok=True)
    try:
        import pymupdf
        doc = pymupdf.open(args.pptx_pdf)
        if doc.page_count < 1:
            print("PPTX→PDF 是 0 页：%s" % args.pptx_pdf); return 2
        pptx_png = os.path.join(args.work, "pptx_render.png")
        doc[0].get_pixmap(dpi=int(round(args.dpi))).save(pptx_png)
    except Exception as exc:
        print("读 %s 失败：%s" % (args.pptx_pdf, exc)); return 2
    target_mm = None
    if args.pdf_size:
        try:
            target_mm = parse_size(args.pdf_size)
        except ValueError as exc:
            print("! %s" % exc); return 1

    H = np.asarray(Image.open(args.pdf_preview).convert("RGB")).astype(np.int16)
    P = np.asarray(Image.open(pptx_png).convert("RGB")).astype(np.int16)
    shapes = json.load(open(args.shapes, encoding="utf-8"))
    canvas_w = args.canvas_w or float((shapes.get("canvas") or {}).get("w") or 3402)
    hs, ps = H.shape[1] / canvas_w, P.shape[1] / canvas_w

    ok = True
    print("1) 编号居中（徽标内数字，白色墨迹中心）")
    dxs, dys = [], []
    hv = [t for t in shapes["texts"] if t.get("mode") == "hv"]   # 编号 = hv 文本项，与类名解耦
    no_ink = 0
    for tx in hv:
        b = tx["box"]
        cx, cy = b["x"] + b["w"] / 2, b["y"] + b["h"] / 2
        a, p = ink_box(H, hs, b), ink_box(P, ps, b)
        if not a or not p:
            print("   %-20s 无墨迹 html=%s pptx=%s" % (tx["sel"], bool(a), bool(p)))
            no_ink += 1
            continue
        dx = ((p[0] + p[2]) / 2 / ps - cx) - ((a[0] + a[2]) / 2 / hs - cx)
        dy = ((p[1] + p[3]) / 2 / ps - cy) - ((a[1] + a[3]) / 2 / hs - cy)
        dxs.append(dx); dys.append(dy)
        if max(abs(dx), abs(dy)) > args.tol_center:
            print("   %-20s 偏移 dX=%+.2f dY=%+.2f" % (tx["sel"], dx, dy))
    if hv and dxs and not no_ink:
        worst = max(max(abs(v) for v in dxs), max(abs(v) for v in dys))
        print("   badges %d | dX mean %+.2f  dY mean %+.2f  max|.| %.2f px  (tol %.2f)"
              % (len(dxs), np.mean(dxs), np.mean(dys), worst, args.tol_center))
        ok = ok and worst <= args.tol_center
    elif args.no_badge_check:
        print("   [skip] 没有可检查的编号徽标（--no-badge-check 已显式放行）")
    else:
        # "没有可检查的编号"不能算通过：要么这份墙报真的有编号却一个都没采集到，
        # 要么编号方块在 PPTX 里是空的 —— 两种都必须看见
        print("   FAIL: 编号检查没能跑（实测 hv 文本项 %d 个、有墨迹可比对 %d 个、"
              "无墨迹 %d 个）；若这份墙报确实没有编号徽标，请加 --no-badge-check"
              % (len(hv), len(dxs), no_ink))
        ok = False

    print("2) 文本落位（墨迹顶边差，PPT − HTML 预览）")
    rows = []
    for tx in shapes["texts"]:
        b = tx["box"]
        win = {"x": b["x"], "y": b["y"], "w": b["w"], "h": b["h"]}
        a, p = ink_box(H, hs, win), ink_box(P, ps, win)
        if a and p:
            rows.append((abs(p[1] / ps - a[1] / hs), p[1] / ps - a[1] / hs, tx["sel"], tx.get("para", 0)))
    if rows:
        d = np.array([r[1] for r in rows])
        worst = max(r[0] for r in rows)
        print("   items %d | dTop mean %+.2f  median %+.2f  p90 %.2f  max %.2f px  (tol %.1f)"
              % (len(rows), d.mean(), np.median(d), np.percentile(np.abs(d), 90), worst, args.tol_top))
        for r in sorted(rows, reverse=True)[:5]:
            if r[0] > 3:
                print("      %-26s para%-2d %+.1f px" % (r[2], r[3], r[1]))
        ok = ok and worst <= args.tol_top
    else:
        print("   ! 一个文本项都没比对到：检查 --pdf-preview / --pptx-pdf 是否是同一版画布")
        ok = False

    print("3) 图片（变形 / 裁切 / 缺失）")
    # 源图根目录：优先 --html 所在目录；没给就按 shapes.json 所在目录（与 20_ 的解析口径一致）
    roots = []
    if args.html:
        roots.append(os.path.dirname(os.path.abspath(args.html)))
    else:
        roots.append(os.path.dirname(os.path.abspath(args.shapes)))
        print("   note: 未给 --html，源图按 %s 解析" % roots[-1])
    pics, read_err = [], None
    try:
        from pptx import Presentation
        pics = [sh for sh in Presentation(args.pptx).slides[0].shapes if sh.shape_type == PICTURE]
    except Exception as exc:
        read_err = str(exc)[:120]
        print("   FAIL: 读不出 PPTX 里的图片（%s）—— 这份 PPTX 不能算验收通过" % read_err)
    # 按 20_html2pptx.py 的命名规则把"实测图片"与"PPTX 图片"对上（不能靠顺序：
    # 任何一张图被跳过都会让后面全部错位，比较的就不是同一张图了）
    seq, expect_name = {}, []
    for im in shapes["images"]:
        b = os.path.basename(str(im.get("src") or "image").replace("/", os.sep))
        seq[b] = seq.get(b, 0) + 1
        expect_name.append("FIG %s%s" % (b, "" if seq[b] == 1 else "#%d" % seq[b]))
    by_name = {}
    for sh in pics:
        by_name.setdefault(sh.name, []).append(sh)
    used, absent = 0, []
    bad = missing = cover = stretch = 0
    for idx, im in enumerate(shapes["images"]):
        b = im["box"]
        rel = im["src"].replace("/", os.sep) if im.get("src") else ""
        p = next((os.path.normpath(os.path.join(r, rel)) for r in roots
                  if rel and os.path.isfile(os.path.normpath(os.path.join(r, rel)))), None)
        if not p:
            print("   %-40s 找不到源图（无法比对；用 --html 指定根目录）"
                  % os.path.basename(str(im.get("src") or "?")))
            missing += 1
            continue
        try:
            src = Image.open(p)
            src_ar = src.width / src.height
        except Exception as exc:
            print("   %-40s 源图打不开（%s）" % (os.path.basename(p), str(exc)[:40]))
            missing += 1
            continue
        fit = (im.get("fit") or "fill").lower()
        pic = (by_name.get(expect_name[idx]) or [None])[0]
        if pic is None:                      # 名字对不上就退回顺序配对，但要说清楚
            pic = pics[idx] if idx < len(pics) else None
            if pic is not None:
                print("   %-40s PPTX 里没找到对象 %r（按顺序配对）"
                      % (os.path.basename(p), expect_name[idx]))
        if pic is None:
            absent.append(expect_name[idx])
            continue
        used += 1
        if fit in ("cover", "contain", "scale-down"):
            cover += 1
            frame_ar = pic.width / pic.height
            cropped = any((pic.crop_left, pic.crop_top, pic.crop_right, pic.crop_bottom))
            if fit == "cover":
                if abs(frame_ar / src_ar - 1) > 0.005 and not cropped:
                    print("   %-40s object-fit:cover 但 PPTX 里没有裁切 —— 横幅会被拉伸"
                          % os.path.basename(p))
                    stretch += 1
            elif abs(frame_ar / src_ar - 1) > 0.005:
                print("   %-40s object-fit:%s 的图片框比例 %.4f ≠ 源图 %.4f（应等比）"
                      % (os.path.basename(p), fit, frame_ar, src_ar))
                stretch += 1
            continue
        if im.get("tf"):
            continue                       # 旋转图：rect 是旋转后包围盒，比例天然不等（20_ 用未变换尺寸画）
        if abs((b["w"] / b["h"]) / src_ar - 1) * 100 > args.aspect_tol:
            print("   %-40s 实测框 Δ%.2f%%（HTML 写的 w/h 与源图不符）"
                  % (os.path.basename(p), abs((b["w"] / b["h"]) / src_ar - 1) * 100)); bad += 1
    print("   distorted %d | 源图缺失 %d | 裁切/等比图 %d | 应裁未裁 %d | PPTX 里没有对应图 %d"
          "（PPTX 图片对象 %d / 实测图片 %d，容差 %.2f%%）"
          % (bad, missing, cover, stretch, len(absent), len(pics), len(shapes["images"]),
             args.aspect_tol))
    ok = ok and bad == 0 and missing == 0 and stretch == 0 and read_err is None
    if absent and not args.allow_missing_pictures:
        print("   FAIL: 实测到 %d 张图，PPTX 里只有 %d 个图片对象（缺 %s）—— 缺图的 PPTX 不算通过"
              % (len(shapes["images"]), used, absent[:3]))
        ok = False
    elif absent:
        print("   [skip] 缺 %d 张图（--allow-missing-pictures 已显式放行）" % len(absent))

    print("4) 包结构 / 页尺寸")
    probs = check_package(args.pptx)
    for msg in probs[:8]:
        print("   %s" % msg)
    n_parts = 0
    try:                                   # 包不合法时**不要**再构造一次 ZipFile（会 BadZipFile 崩）
        with zipfile.ZipFile(args.pptx) as z:
            n_parts = len(z.namelist())
    except Exception as exc:
        probs.append("不是合法 zip：%s" % exc)
    print("   package: %s（%d 个部件）" % ("pass" if not probs else "FAIL", n_parts))
    ok = ok and not probs
    r = doc[0].rect
    w_mm, h_mm = r.width / 72 * 25.4, r.height / 72 * 25.4
    if target_mm:
        tw, th = target_mm
    else:
        c = shapes.get("canvas") or {}
        tw, th = c.get("w", 0) / 96 * 25.4, c.get("h", 0) / 96 * 25.4
    print("   pptx->pdf: %d page(s) %.2f x %.2f mm (target %.2f x %.2f)"
          % (len(doc), w_mm, h_mm, tw, th))
    ok = ok and len(doc) == 1 and abs(w_mm - tw) <= 1.0 and abs(h_mm - th) <= 1.0

    print("VERIFY: %s" % ("PASS" if ok else "FAIL"))
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main())
