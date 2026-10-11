# -*- coding: utf-8 -*-
"""20_html2pptx.py —— 由 DOM 实测 JSON 生成可编辑 PPTX（不需要浏览器）。

    python 20_html2pptx.py --shapes _work/shapes.json --out poster.pptx \
        --html poster.html --font "Microsoft YaHei" --slide-mm 900x1200

规则见 references/02-pptx-conversion.md；本文件是那些规则的参考实现：
  内容盒定位 / 按段拆文本框 / 精确行距 / 底对齐修正 / 居中文案对称加宽 / 编号双向居中 /
  圆角 adj=radius/min(w,h) / 四条边框（含单边与虚线）/ object-fit 裁切 / 图片圆角与旋转 /
  下划线·删除线·大小写变换 / 上下标按实测基线偏移 / 半透明色按背景合成 / 空心项目符号 /
  色值全部来自实测，不写死主题色。

退出码：0 = 正常；1 = 有图片缺失（默认不产出残缺 PPTX，除非 --allow-missing-images）。
"""
from __future__ import annotations

import argparse
import json
import os
import pathlib
import re
import sys
import urllib.parse

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

for _s in (sys.stdout, sys.stderr):        # Windows 中文控制台默认 GBK：打印 "−/≈/✓" 会崩
    try:
        _s.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

NS = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def parse_pos(txt: str):
    """CSS object-position → (横向 token, 纵向 token)。"""
    toks = str(txt or "50% 50%").replace(",", " ").split()
    if not toks:
        return "50%", "50%"
    if len(toks) == 1:
        return toks[0], toks[0] if toks[0] in ("top", "bottom") else "50%"
    return toks[0], toks[1]


def pos_frac(tok: str, avail: float) -> float:
    """token → [0,1] 定位比例。avail = 该轴可滑动范围（px）；px 值按 CSS 语义换算。"""
    tok = (tok or "").strip().lower()
    if tok.endswith("%"):
        return float(tok[:-1]) / 100.0
    if tok.endswith("px"):
        return 0.0 if avail <= 0 else min(1.0, max(0.0, float(tok[:-2]) / avail))
    return {"left": 0.0, "top": 0.0, "center": 0.5, "right": 1.0, "bottom": 1.0}.get(tok, 0.5)


def transform_text(t: str, tt: str) -> str:
    """CSS text-transform 在浏览器里是渲染期变换：PPT 里要落到文字上。"""
    if tt == "uppercase":
        return t.upper()
    if tt == "lowercase":
        return t.lower()
    if tt == "capitalize":
        return re.sub(r"(^|\s)(\S)", lambda m: m.group(1) + m.group(2).upper(), t)
    return t


def place(box, tf, what, warns):
    """返回 (x, y, w, h, rot)。纯旋转用未变换尺寸 + 旋转角；缩放/斜切只告警。"""
    if not tf:
        return box["x"], box["y"], box["w"], box["h"], 0
    if tf.get("other"):
        warns.append("有 scale/skew 变换，PPT 里按未变换框绘制：%s" % what)
        return box["x"], box["y"], box["w"], box["h"], 0
    ow, oh = tf.get("ow") or box["w"], tf.get("oh") or box["h"]
    return (box["x"] + box["w"] / 2 - ow / 2, box["y"] + box["h"] / 2 - oh / 2,
            ow, oh, tf.get("rot") or 0)


def round_picture(pic, radius_px: float, w: float, h: float) -> None:
    """把图片几何改成圆角矩形（等价 PPT"裁剪为形状"），对应 CSS border-radius。"""
    spPr = pic._element.spPr
    prst = spPr.find(qn("a:prstGeom"))
    if prst is None:
        return
    prst.set("prst", "roundRect")
    av = prst.find(qn("a:avLst"))
    if av is None:
        av = prst.makeelement(qn("a:avLst"), {})
        prst.append(av)
    for ch in list(av):
        av.remove(ch)
    adj = int(round(min(0.5, radius_px / max(1.0, min(w, h))) * 100000))
    av.append(av.makeelement(qn("a:gd"), {"name": "adj", "fmla": "val %d" % adj}))


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--shapes", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--html", default=None, help="仅用于解析图片相对路径")
    ap.add_argument("--font", default="Microsoft YaHei")
    ap.add_argument("--slide-mm", default="900x1200")
    ap.add_argument("--emu-per-px", type=int, default=9525)      # 96 dpi
    ap.add_argument("--px-to-pt", type=float, default=0.75)
    ap.add_argument("--allow-missing-images", action="store_true",
                    help="有图片缺失时也照常产出（默认视为失败，避免交出缺图的 PPTX）")
    ap.add_argument("--allow-distorted-images", action="store_true",
                    help="图片与框纵横比不符（会被拉伸）时也照常产出；默认判失败")
    args = ap.parse_args()

    E = lambda v: Emu(int(round(v * args.emu_per_px)))
    # --html 缺省时按 shapes.json 所在目录解析图片（否则相对 src 会被拿 CWD 去拼）
    if not args.html:
        print("note: 未给 --html，图片相对路径按 %s 解析"
              % os.path.dirname(os.path.abspath(args.shapes)))
    base = (os.path.dirname(os.path.abspath(args.html)) if args.html
            else os.path.dirname(os.path.abspath(args.shapes)))
    try:
        with open(args.shapes, encoding="utf-8") as fh:
            data = json.load(fh)
    except FileNotFoundError:
        print("! 找不到实测数据（--shapes）：%s" % os.path.abspath(args.shapes))
        return 2
    except json.JSONDecodeError as exc:
        print("! %s 不是合法 JSON（行 %d）：%s" % (args.shapes, exc.lineno, exc.msg))
        print("  提示：shapes.json 由 10_measure_gate_export.py 生成；被截断或改坏时重新生成")
        return 1
    except Exception as exc:
        print("! 读 %s 失败：%s" % (args.shapes, exc))
        return 2
    if not isinstance(data, dict) or not data.get("texts") and not data.get("images") and not data.get("rects"):
        print("! %s 里没有 rects/texts/images —— 不是本工具产出的实测数据" % args.shapes)
        return 1
    os.makedirs(os.path.dirname(os.path.abspath(args.out)) or ".", exist_ok=True)
    try:
        sw, sh = (float(x) for x in args.slide_mm.lower().split("x"))    # 毫米
    except Exception:
        print("! --slide-mm 要写成 WxH 毫米，例如 900x1200；收到：%r" % args.slide_mm)
        return 1
    EMU_PER_MM = 36000
    warns = []
    seq = {}          # 名称去重：同一前缀的元素加 #n 后缀
    # 字体：本机没有指定字体时必须告警（PPT 会回退到别的字体 → 字宽不同 → 断行点与 PDF 不一致）
    try:
        import importlib.util as _ilu
        _spec = _ilu.spec_from_file_location(
            "_apb_check_env", str(pathlib.Path(__file__).resolve().parent / "00_check_env.py"))
        _ce = _ilu.module_from_spec(_spec)
        _spec.loader.exec_module(_ce)
        if not _ce.find_font(args.font):
            warns.append("字体 %r 在本机没找到 —— PowerPoint 会回退到别的字体，断行点可能与浏览器不同；"
                         "用 --font 指定本机实际安装的字体族" % args.font)
    except Exception:
        pass

    prs = Presentation()
    prs.slide_width = Emu(int(sw * EMU_PER_MM))
    prs.slide_height = Emu(int(sh * EMU_PER_MM))
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    shapes = slide.shapes

    canvas = data.get("canvas") or {}
    if canvas.get("w"):
        mm_w, mm_h = canvas["w"] / 96 * 25.4, canvas["h"] / 96 * 25.4
        if abs(mm_w - sw) > sw * 0.005 or abs(mm_h - sh) > sh * 0.005:
            warns.append("画布 %gx%g px = %.1fx%.1f mm，但 --slide-mm %s；内容会被裁切"
                         % (canvas["w"], canvas["h"], mm_w, mm_h, args.slide_mm))

    def set_face(run, face):
        """latin + east-asian + complex-script 三处同时设族（否则中文走别的回退字体）"""
        rPr = run._r.get_or_add_rPr()
        for tag in ("latin", "ea", "cs"):
            el = rPr.find(NS + tag)
            if el is None:
                el = rPr.makeelement(NS + tag, {}); rPr.append(el)
            el.set("typeface", face)

    def style_run(r, run):
        """一个 run 的全部行内属性：字号/粗斜/字体/颜色/下划线/删除线/上下标。"""
        f = r.font
        f.size = Pt(round(run["size"] * args.px_to_pt, 1))
        f.bold, f.italic, f.name = bool(run["bold"]), bool(run["italic"]), args.font
        set_face(r, args.font)
        c = run.get("color")
        if c is None or (isinstance(c, int) and c < 0) or run.get("colorOk") is False:
            warns.append("颜色没解析出来（可能是 oklch/lab/color-mix 等色彩空间），已按黑色输出：%s"
                         % str(run.get("t"))[:16])
            c = 0
        f.color.rgb = RGBColor((c >> 16) & 255, (c >> 8) & 255, c & 255)
        if run.get("u"):
            f.underline = True
        if run.get("s"):
            r._r.get_or_add_rPr().set("strike", "sngStrike")
        # 上下标：优先用实测的基线偏移（<sup>/<sub>/自定义 vertical-align 都适用）
        shift = run.get("shift") or 0
        bl = None
        # 物理上限：上下标偏移不会超过 0.5 em；超过就说明测量异常，宁可不设
        if 1.2 < abs(shift) <= 0.5 * run["size"]:
            bl = int(round(-shift / max(1.0, run["size"]) * 100000))
        elif abs(shift) > 0.5 * run["size"]:
            warns.append("run 的基线偏移异常（%.1f px > 0.5em），已忽略：%s" % (shift, run["t"][:16]))
        elif run.get("sub"):
            bl = -25000
        if bl:
            r._r.get_or_add_rPr().set("baseline", str(bl))

    def set_custgeom(shp, box, corners):
        """逐角圆角矩形（custGeom）：对应 CSS border-radius + 父容器 overflow:hidden 的裁切。"""
        from lxml import etree
        W = int(round(box["w"] * args.emu_per_px))
        H = int(round(box["h"] * args.emu_per_px))
        lim = min(box["w"], box["h"]) / 2.0
        tl, tr, br, bl = [int(round(min(c, lim) * args.emu_per_px)) for c in corners]
        spPr = shp._element.spPr
        for tag in ("a:prstGeom", "a:custGeom"):
            e = spPr.find(qn(tag))
            if e is not None:
                spPr.remove(e)
        NS_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
        x = ['<a:custGeom xmlns:a="%s"><a:avLst/><a:gdLst/><a:ahLst/><a:cxnLst/>' % NS_A,
             '<a:rect l="0" t="0" r="%d" b="%d"/>' % (W, H),
             '<a:pathLst><a:path w="%d" h="%d">' % (W, H)]
        pt = lambda px, py: '<a:pt x="%d" y="%d"/>' % (px, py)
        arc = lambda rr, st: '<a:arcTo wR="%d" hR="%d" stAng="%d" swAng="5400000"/>' % (rr, rr, st)
        x.append("<a:moveTo>" + pt(tl, 0) + "</a:moveTo>")
        x.append("<a:lnTo>" + pt(W - tr, 0) + "</a:lnTo>")
        if tr: x.append(arc(tr, 16200000))
        x.append("<a:lnTo>" + pt(W, H - br) + "</a:lnTo>")
        if br: x.append(arc(br, 0))
        x.append("<a:lnTo>" + pt(bl, H) + "</a:lnTo>")
        if bl: x.append(arc(bl, 5400000))
        x.append("<a:lnTo>" + pt(0, tl) + "</a:lnTo>")
        if tl: x.append(arc(tl, 10800000))
        x.append("<a:close/></a:path></a:pathLst></a:custGeom>")
        geom = etree.fromstring("".join(x))
        xfrm = spPr.find(qn("a:xfrm"))
        if xfrm is not None:
            xfrm.addnext(geom)
        else:
            spPr.insert(0, geom)

    def add_rect(box, fill=None, line=None, lw=0.0, rounded=False, radius=0.0, dash=None, rot=0.0,
                 corners=None):
        """fill/line = None 表示"没有颜色"。
        注意：**不能用真值判断**——黑色 (0x000000) 是合法颜色，旧版 `if fill:` 会把黑色填充、
        黑色边框、黑色单边线、黑色项目符号全部静默丢掉（深色配色的墙报会整块变透明）。"""
        if rounded:
            shp = shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                   E(box["x"]), E(box["y"]), E(box["w"]), E(box["h"]))
            try:
                shp.adjustments[0] = min(0.5, radius / max(1.0, min(box["w"], box["h"])))
            except Exception:
                pass
        else:
            shp = shapes.add_shape(MSO_SHAPE.RECTANGLE,
                                   E(box["x"]), E(box["y"]), E(box["w"]), E(box["h"]))
        if fill is not None:
            shp.fill.solid(); shp.fill.fore_color.rgb = RGBColor((fill >> 16) & 255, (fill >> 8) & 255, fill & 255)
        else:
            shp.fill.background()
        if line is not None and lw > 0:
            shp.line.color.rgb = RGBColor((line >> 16) & 255, (line >> 8) & 255, line & 255)
            shp.line.width = Pt(lw * args.px_to_pt)
            if dash:
                from pptx.enum.dml import MSO_LINE_DASH_STYLE
                shp.line.dash_style = MSO_LINE_DASH_STYLE.DASH if dash == "dashed" else MSO_LINE_DASH_STYLE.ROUND_DOT
        else:
            shp.line.fill.background()
        shp.shadow.inherit = False
        if corners:
            set_custgeom(shp, box, corners)
        if rot:
            shp.rotation = rot
        return shp

    def add_text(item):
        box, runs, mode, align = item["box"], item["runs"], item.get("mode", ""), item["align"]
        lhs = item.get("lhpx") or 0
        h1 = item.get("fh") or box["h"]
        # 浏览器：字体盒居中于行盒；PPT/LO 精确行距：字体盒底对齐行盒 → 需要下移 (H − S)
        dy = (h1 - lhs) if (lhs and mode != "hv") else 0.0
        height_slack = 0 if mode == "hv" else 24
        if mode == "hv":                       # 徽标：精确盒 + 双向居中，不加宽、不下移
            wslack, xoff, align_ppt = 0, 0, PP_ALIGN.CENTER
        elif align in ("left", "start"):
            wslack, xoff, align_ppt = 6, 0, PP_ALIGN.LEFT
        elif align == "center":
            wslack, xoff, align_ppt = 48, 24, PP_ALIGN.CENTER      # 两侧对称加宽，中心不动
        elif align in ("right", "end"):
            wslack, xoff, align_ppt = 48, 48, PP_ALIGN.RIGHT       # 向左加宽，右边界不动
        elif align == "justify":
            wslack, xoff, align_ppt = 6, 0, PP_ALIGN.JUSTIFY
        else:                                  # 未知对齐：按左对齐处理并告警，不要默默右对齐
            wslack, xoff, align_ppt = 6, 0, PP_ALIGN.LEFT
            warns.append("未识别的 text-align=%r，已按左对齐输出：%s"
                         % (align, "".join(r["t"] for r in runs)[:16]))
        if item.get("tf"):
            warns.append("文本框有 transform，PPT 里不旋转：%s" % ("".join(r["t"] for r in runs)[:20]))
        if item.get("outline"):
            warns.append("文本元素有 outline，PPT 里不绘制：%s" % ("".join(r["t"] for r in runs)[:16]))
        if item.get("clipped"):
            warns.append("该元素被祖先 overflow:hidden（含圆角）裁掉一部分，PPT 里不会裁切：%s"
                         % ("".join(r["t"] for r in runs)[:18]))
        if item.get("overflowX"):
            warns.append("内容比盒子宽（超长单词/nowrap），PPT 里已关闭自动换行：%s"
                         % ("".join(r["t"] for r in runs)[:20]))
        tb = shapes.add_textbox(E(box["x"] - xoff), E(box["y"] + dy),
                                E(box["w"] + wslack), E(box["h"] + height_slack))
        tf = tb.text_frame
        tf.word_wrap = not item.get("nowrap")   # white-space:nowrap 在 PPT 里也不能换行
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE if mode == "hv" else MSO_ANCHOR.TOP
        p = tf.paragraphs[0]
        p.alignment = align_ppt
        if lhs and mode != "hv":
            p.line_spacing = Pt(round(lhs * args.px_to_pt, 1))     # 精确磅值，不用倍数
        for run in runs:
            tt = run.get("tt") or "none"
            txt = transform_text(run["t"], tt) if tt != "none" else run["t"]
            if "\t" in txt:                    # 制表符在 <a:t> 里无法表达，按 4 空格近似
                txt = txt.replace("\t", "    ")
                warns.append("文本里有制表符，已按 4 空格近似：%s" % run["t"][:16])
            # <pre>/white-space:pre 的换行要是段落内软换行；直接塞 \n 会被转义成 _x000A_
            for k, seg in enumerate(txt.split("\n")):
                if k:
                    p._p.append(p._p.makeelement(qn("a:br"), {}))
                if seg == "" and k == 0 and txt.startswith("\n"):
                    continue
                r = p.add_run()
                r.text = seg
                style_run(r, run)
        first = "".join(x["t"] for x in runs)[:32]
        if item.get("npara", 1) > 1:
            first = "%s [%d/%d]" % (first, item["para"] + 1, item["npara"])
        key = "TXT %s" % first
        seq[key] = seq.get(key, 0) + 1
        tb.name = key if seq[key] == 1 else "%s#%d" % (key, seq[key])
        return tb

    # 1) 形状（面板/卡片底 → 标题条 → 边框/项目符号 → 图片 → 文本，保证 z-order）。
    #    色值全部来自浏览器实测，不写死主题色：写死的那版换配色会静默错版。
    #    实测里的 -1 = "没有颜色"（与黑色 0 区分开）。
    def col(v):
        return None if v is None or (isinstance(v, int) and v < 0) else v

    for rc in data["rects"]:
        box, radius, fill = rc["box"], rc.get("radius") or 0, col(rc.get("fill"))
        bw = rc.get("bw") or [0, 0, 0, 0]                  # 上右下左
        bc = [col(c) for c in (rc.get("bc") or [0, 0, 0, 0])]
        rnd = radius > 0                                   # radius==0 用普通矩形（roundRect 有渲染差）
        x, y, w, h, rot = place(box, rc.get("tf"), "形状 %s" % rc.get("sel"), warns)
        box = {"x": x, "y": y, "w": w, "h": h}
        dash = rc.get("bstyle") if rc.get("bstyle") in ("dashed", "dotted") else None
        if rc.get("shadow"):
            warns.append("元素有 box-shadow，PPT 里不绘制：%s" % rc.get("sel"))
        if rc.get("clipped"):
            warns.append("该形状被祖先 overflow:hidden（含圆角）裁掉一部分，PPT 里不会裁切：%s" % rc.get("sel"))
        if rc.get("outline"):
            warns.append("元素有 outline（浏览器会画，PPT 形状没有这一层）：%s" % rc.get("sel"))
        if rc.get("corners"):
            pass          # 已按父容器圆角逐角裁切，不需要告警
        elif rc.get("clipRisk"):
            warns.append("父容器圆角 + overflow:hidden 会裁切子元素背景，PPT 不裁切：%s" % rc.get("sel"))
        if any(st in ("double", "groove", "ridge", "inset", "outset") for st in (rc.get("bstyles") or [])):
            warns.append("边框是 double/groove/ridge/inset/outset，PPT 里按实线画：%s" % rc.get("sel"))
        if bw[0] > 0 and bw[0] == bw[1] == bw[2] == bw[3]:
            add_rect(box, fill, line=bc[0], lw=bw[0], rounded=rnd and not rc.get("corners"),
                     radius=radius, dash=dash, rot=rot, corners=rc.get("corners"))
            continue
        add_rect(box, fill, rounded=rnd and not rc.get("corners"), radius=radius, rot=rot,
                 corners=rc.get("corners"))
        for sx, sy, sw_, sh_, cl in ((x, y, w, bw[0], bc[0]),                 # 上
                                     (x + w - bw[1], y, bw[1], h, bc[1]),     # 右
                                     (x, y + h - bw[2], w, bw[2], bc[2]),     # 下
                                     (x, y, bw[3], h, bc[3])):                # 左
            if sw_ > 0 and sh_ > 0 and cl is not None:     # 单边边框（.panel-head 下边、.note 左边）
                add_rect({"x": sx, "y": sy, "w": sw_, "h": sh_}, cl)
    for bl in data.get("bullets", []):         # 项目符号：几何/颜色/描边都来自 ::before
        bf, bl_line = col(bl.get("fill")), col(bl.get("line"))
        if bf is not None or bl_line is not None:
            br = bl.get("radius") or 0
            add_rect(bl["box"], bf, line=bl_line, lw=bl.get("lw") or 0, rounded=br > 0, radius=br)
    if data.get("bulletsSkipped"):
        print("  · %d 个 li 用原生/文字项目符号，PPT 里不会复现（要复现就改成有尺寸的 ::before 方块）"
              % data["bulletsSkipped"])

    # 2) 图片：cover→按 srcRect 裁切（直接拉伸会把横幅压扁），contain→等比缩放后按 object-position 摆位
    missing = 0
    unsupported = 0
    seen = {}
    warned_src = set()

    def img_path(src):
        """src → 本地文件路径；返回 (path 或 None, 原因)。
        data:/http(s):/file:// 不是本地相对路径，不能当"文件缺失"处理
        （旧版会把它们拼成 D:\\proj\\data:image\\... 然后报缺失）。"""
        if not src:
            return None, "无 src（懒加载 data-src / 只用 srcset 的写法）"
        s = str(src).strip()
        low = s.lower()
        if low.startswith("data:"):
            return None, "data: URI（内嵌图，OOXML 生成器不解析）"
        if low.startswith(("http://", "https://", "//")):
            return None, "远程地址（PPTX 只能内嵌本地图片）"
        if low.startswith("file://"):
            return pathlib.Path(urllib.parse.unquote(s[7:].lstrip("/"))), ""
        s = urllib.parse.unquote(s.split("#")[0].split("?")[0])          # 去掉 ?v= / #frag
        return os.path.normpath(os.path.join(base, s.replace("/", os.sep))), ""

    for im in data["images"]:
        path, why = img_path(im.get("src"))
        if path is None:
            print("  ! 图片跳过（%s）：%s" % (why, str(im.get("src"))[:60]))
            unsupported += 1
            continue
        if not os.path.isfile(path):
            print("  ! 图片缺失：%s" % path)
            missing += 1
            continue
        from PIL import Image as _I
        try:
            with _I.open(path) as src:
                iw, ih = src.size
                src.verify()                    # 0 字节/截断/改名文件在这里暴露，而不是 traceback
        except Exception as exc:
            print("  ! 图片不是可读的位图（%s），已跳过：%s" % (str(exc).splitlines()[0][:60], path))
            unsupported += 1
            continue
        box = dict(im["box"])
        fit = (im.get("fit") or "fill").lower()
        px_tok, py_tok = parse_pos(im.get("pos"))
        crop = (0.0, 0.0, 0.0, 0.0)
        if fit in ("cover", "contain", "scale-down"):
            if fit == "cover":
                sc = max(box["w"] / iw, box["h"] / ih)
                vw, vh = box["w"] / sc, box["h"] / sc
                fx, fy = pos_frac(px_tok, iw - vw), pos_frac(py_tok, ih - vh)
                sx, sy = (iw - vw) * fx, (ih - vh) * fy
                crop = (sx / iw, sy / ih, 1 - (sx + vw) / iw, 1 - (sy + vh) / ih)
            else:
                sc = min(box["w"] / iw, box["h"] / ih)
                if fit == "scale-down":
                    sc = min(sc, 1.0)
                dw, dh = iw * sc, ih * sc
                box["x"] += (box["w"] - dw) * pos_frac(px_tok, box["w"] - dw)
                box["y"] += (box["h"] - dh) * pos_frac(py_tok, box["h"] - dh)
                box["w"], box["h"] = dw, dh
        x, y, w, h, rot = place(box, im.get("tf"), "图片 %s" % str(im.get("src"))[:40], warns)
        pic = shapes.add_picture(path, E(x), E(y), E(w), E(h))
        if im.get("clipped"):
            warns.append("该图片被祖先 overflow:hidden（含圆角）裁掉一部分，PPT 里不会裁切：%s" % im.get("src"))
        if any(crop):
            pic.crop_left, pic.crop_top, pic.crop_right, pic.crop_bottom = crop
        if im.get("radius"):
            round_picture(pic, im["radius"], w, h)
        if rot:
            pic.rotation = rot
        b = os.path.basename(str(im.get("src") or "image").replace("/", os.sep))
        seen[b] = seen.get(b, 0) + 1
        pic.name = "FIG %s%s" % (b, "" if seen[b] == 1 else "#%d" % seen[b])

    # 3) 文本
    for tx in data["texts"]:
        add_text(tx)

    # 4) 产出前自检：图片变形、元素越界（**必须报出来，并且默认判失败**）
    problems = 0
    distorted = []
    cw, ch = canvas.get("w", 0), canvas.get("h", 0)
    for im in data["images"]:
        b = im["box"]
        path, _why = img_path(im.get("src"))
        if path is None or not os.path.isfile(path):
            continue
        if (im.get("fit") or "fill") in ("cover", "contain", "scale-down"):
            continue                       # 裁切/等比缩放：显示框比例本来就不等于源图
        if im.get("tf"):
            continue                       # 旋转过的图：rect 是**旋转后包围盒**，比例天然不等
        from PIL import Image as _I
        try:
            with _I.open(path) as src:
                dev = abs((b["w"] / b["h"]) / (src.width / src.height) - 1) * 100
        except Exception:
            continue
        if dev > 0.5:
            print("  ! 图片变形 %.2f%%：%s（HTML %gx%g）" % (dev, im.get("src"), b["w"], b["h"]))
            distorted.append(im.get("src"))
            problems += 1
    if cw and ch:
        for item in data["texts"] + data["rects"]:
            b = item["box"]
            if b["x"] < -2 or b["y"] < -2 or b["x"] + b["w"] > cw + 2 or b["y"] + b["h"] > ch + 2:
                print("  ! 越界：%s %s" % (item.get("sel", "?"), {k: round(v) for k, v in b.items()}))
                problems += 1
    if (data.get("nestedCols") or 0):
        warns.append("检测到 %d 个嵌套 .col（.col 里还有 .col）：栏底对齐/栏区边距会失真，"
                     "请把列改成同级" % data["nestedCols"])
    for w in sorted(set(warns)):
        print("  · %s" % w)
    if problems:
        print("  自检：%d 处问题（见上）" % problems)
    if missing and not args.allow_missing_images:
        print("  ! %d 张图片没找到 —— 缺图的 PPTX 不是可交付版本；"
              "修好路径，或明确加 --allow-missing-images" % missing)
        return 1
    if unsupported and not args.allow_missing_images:
        print("  ! %d 张图片无法内嵌（data:/远程/SVG/坏文件）—— 修好或显式加 "
              "--allow-missing-images" % unsupported)
        return 1
    if distorted and not args.allow_distorted_images:
        print("  ! %d 张图片与显示框纵横比不符（会被拉伸）——PPT 里会明显变形；"
              "改 width/height，或显式加 --allow-distorted-images" % len(distorted))
        return 1

    prs.save(args.out)
    print("written: %s (%.2f MB) | %d shapes | slide %gx%g mm"
          % (args.out, os.path.getsize(args.out) / 2 ** 20, len(shapes), sw, sh))
    return 0


if __name__ == "__main__":
    sys.exit(main())
